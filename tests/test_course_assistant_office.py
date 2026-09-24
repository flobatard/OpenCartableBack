"""Extraction du texte des pièces jointes bureautiques (docx, pptx, odt, odp).

Les documents sont construits EN MÉMOIRE par les libs elles-mêmes : aucun
fichier d'exemple à maintenir, et le test décrit ce qu'on promet au modèle —
le texte, pas la mise en forme.

Les tests de refus (archive gonflée, entité XML, fichier qui n'est pas une
archive) fabriquent le zip à la main : ce sont eux qui gardent les gardes.
"""

import io
import zipfile

import pytest

from app.course_assistant.office import (
    _EXTRACTORS,
    DOCX_MIME,
    ODP_MIME,
    ODT_MIME,
    OFFICE_MAX_CHARS,
    OFFICE_TOTAL_MAX_BYTES,
    PPTX_MIME,
    extract_office_text,
)


def _docx(paragraphs: list[str], table: list[list[str]] | None = None) -> io.BytesIO:
    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        added = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for cell_index, value in enumerate(row):
                added.cell(row_index, cell_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def _pptx(slides: list[str]) -> io.BytesIO:
    from pptx import Presentation

    presentation = Presentation()
    layout = presentation.slide_layouts[5]  # « Titre seul »
    for title in slides:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
    buffer = io.BytesIO()
    presentation.save(buffer)
    buffer.seek(0)
    return buffer


def _odt(paragraphs: list[str]) -> io.BytesIO:
    from odf.opendocument import OpenDocumentText
    from odf.text import P

    document = OpenDocumentText()
    for paragraph in paragraphs:
        document.text.addElement(P(text=paragraph))
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


# ─────────────────────────────────────────────
# Ce qu'on rend au modèle
# ─────────────────────────────────────────────


def test_docx_paragraphs_are_extracted_in_order():
    content = extract_office_text(_docx(["Chapitre 1", "Le théorème dit ceci."]), DOCX_MIME)
    assert "Chapitre 1" in content
    assert content.index("Chapitre 1") < content.index("Le théorème dit ceci.")


def test_docx_tables_are_flattened_row_by_row():
    """Les cellules sont aplaties : c'est la perte assumée, et l'en-tête le dit."""
    content = extract_office_text(
        _docx(["Barème"], table=[["Question", "Points"], ["Q1", "4"]]), DOCX_MIME
    )
    assert "Question | Points" in content
    assert "Q1 | 4" in content


def test_the_header_warns_the_model_that_the_extraction_is_degraded():
    content = extract_office_text(_docx(["Texte"]), DOCX_MIME)
    assert content.startswith("[Texte brut extrait du document")
    assert "mise en forme" in content


def test_pptx_titles_each_slide():
    """Sans repère, plusieurs diapositives se lisent comme un seul bloc."""
    content = extract_office_text(_pptx(["Les fractions", "Exercices"]), PPTX_MIME)
    assert "## Diapositive 1" in content and "Les fractions" in content
    assert "## Diapositive 2" in content and "Exercices" in content


def test_odt_paragraphs_are_extracted():
    content = extract_office_text(_odt(["Première ligne", "Seconde ligne"]), ODT_MIME)
    assert "Première ligne" in content and "Seconde ligne" in content


def test_odp_is_wired_to_the_same_reader_as_odt():
    """Un .odp vide est bien plus lourd à fabriquer qu'à lire (page maîtresse,
    cadre, zone de texte) : ce qui compte est qu'il emprunte le MÊME
    extracteur — les paragraphes d'un ODF se lisent pareil dans les deux."""
    assert _EXTRACTORS[ODP_MIME] is _EXTRACTORS[ODT_MIME]


def test_an_empty_document_yields_an_empty_string():
    """L'appelant en fait une erreur actionnable (« pas de texte extractible »)."""
    assert extract_office_text(_docx([]), DOCX_MIME) == ""


def test_long_documents_are_truncated_with_a_marker():
    long_paragraph = "a" * (OFFICE_MAX_CHARS + 500)
    content = extract_office_text(_docx([long_paragraph]), DOCX_MIME)
    assert "[Document tronqué" in content
    assert len(content) < len(long_paragraph) + 500


# ─────────────────────────────────────────────
# Gardes
# ─────────────────────────────────────────────


def test_an_unknown_office_mime_is_refused():
    with pytest.raises(ValueError):
        extract_office_text(_docx(["x"]), "application/vnd.ms-excel")


def test_a_file_that_is_not_an_archive_is_refused():
    with pytest.raises(zipfile.BadZipFile):
        extract_office_text(io.BytesIO(b"ceci n'est pas un zip"), DOCX_MIME)


def test_an_archive_that_inflates_beyond_the_cap_is_refused():
    """Une archive de quelques kilo-octets peut décompresser en gigaoctets :
    on borne la taille DÉCLARÉE avant de laisser une lib l'ouvrir."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (OFFICE_TOTAL_MAX_BYTES + 1))
    buffer.seek(0)
    with pytest.raises(ValueError, match="décompressée"):
        extract_office_text(buffer, DOCX_MIME)


def test_an_xml_part_with_an_entity_declaration_is_refused():
    """Ni OOXML ni ODF ne portent de DOCTYPE : le refuser ferme billion-laughs
    sans dépendre de la politique d'une lib tierce."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "word/document.xml",
            b'<?xml version="1.0"?><!DOCTYPE lol [<!ENTITY a "aa">]><w:document/>',
        )
    buffer.seek(0)
    with pytest.raises(ValueError, match="entité"):
        extract_office_text(buffer, DOCX_MIME)
